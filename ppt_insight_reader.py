from __future__ import annotations

from io import BytesIO
from typing import Iterable

import pandas as pd
from pptx import Presentation


def clean_text(value: object) -> str:
    return " ".join(str(value or "").replace("\n", " ").split()).strip()


def iter_shape_text(shape) -> Iterable[str]:
    if hasattr(shape, "text") and clean_text(shape.text):
        yield clean_text(shape.text)

    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [clean_text(cell.text) for cell in row.cells if clean_text(cell.text)]
            if cells:
                yield " | ".join(cells)

    if getattr(shape, "has_chart", False):
        chart = shape.chart
        if chart.has_title:
            title = clean_text(chart.chart_title.text_frame.text)
            if title:
                yield title
        for series in chart.series:
            name = clean_text(getattr(series, "name", ""))
            if name:
                yield f"Chart series: {name}"

    if getattr(shape, "shape_type", None) == 6:
        for grouped_shape in shape.shapes:
            yield from iter_shape_text(grouped_shape)


def slide_title_and_body(texts: list[str]) -> tuple[str, str]:
    if not texts:
        return "PowerPoint slide", ""
    title = texts[0]
    body_parts = []
    for text in texts[1:]:
        if text and text not in body_parts:
            body_parts.append(text)
    if not body_parts:
        body_parts = [title]
    return title, " ".join(body_parts)


def read_pptx_insights(uploaded_file) -> pd.DataFrame:
    uploaded_file.seek(0)
    prs = Presentation(BytesIO(uploaded_file.read()))
    rows = []

    for slide_index, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            for text in iter_shape_text(shape):
                if text and text not in texts:
                    texts.append(text)

        title, body = slide_title_and_body(texts)
        if len(body.split()) < 4:
            continue

        rows.append(
            {
                "insight_id": f"PPT-S{slide_index:02d}",
                "theme": title[:90],
                "insight_text": body,
                "evidence_note": f"Extracted from uploaded PowerPoint slide {slide_index}; verify chart values and source tables before client use.",
            }
        )

    return pd.DataFrame(rows)
